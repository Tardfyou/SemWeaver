"""Build two editable research diagrams from shared geometric primitives.

No style-skill assets, model calls or experimental execution are used.
Dependencies: reportlab==4.4.4, python-pptx==1.0.2; Liberation Sans fonts.
All coordinates and font sizes are PDF points, with a top-left origin.
"""
from pathlib import Path
import argparse
import json
import math
import xml.etree.ElementTree as ET

from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas as pdfcanvas
from reportlab.lib.colors import HexColor
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.text import MSO_ANCHOR
from pptx.util import Pt
from pptx.oxml.xmlchemy import OxmlElement

INK = '#243448'
MUTED = '#617084'
LINE = '#B9C3CF'
PALE = '#F3F5F8'
BLUE = '#315F98'
BLUE_PALE = '#EEF4FC'
AMBER = '#A66024'
AMBER_PALE = '#FFF5E9'
GREEN = '#267560'
GREEN_PALE = '#EDF7F2'
WHITE = '#FFFFFF'


class Figure:
    def __init__(self, name, height):
        self.name, self.width, self.height = name, 378, height
        self.items = []

    def box(self, x, y, w, h, fill=WHITE, stroke=LINE, radius=3, dash=False, lw=.8):
        self.items.append(dict(kind='box', x=x, y=y, w=w, h=h, fill=fill,
                               stroke=stroke, radius=radius, dash=dash, lw=lw))

    def text(self, x, y, value, size=8.5, bold=False, color=INK, center=False, maxw=None):
        font = 'FigureBold' if bold else 'FigureSans'
        width = pdfmetrics.stringWidth(value, font, size)
        if maxw is not None:
            assert width <= maxw, (self.name, value, width, maxw)
        left = x - width / 2 if center else x
        ascent, descent = pdfmetrics.getAscentDescent(font, size)
        assert 0 <= left and left + width <= self.width
        assert y - ascent >= 0 and y - descent <= self.height
        self.items.append(dict(kind='text', x=left, y=y, value=value, size=size,
                               bold=bold, color=color, w=width,
                               ascent=ascent, descent=descent))

    def route(self, points, color=INK, arrow=True, dash=False, lw=1):
        self.items.append(dict(kind='line', points=points, color=color,
                               arrow=arrow, dash=dash, lw=lw))

    def centered(self, x, y, text, **kwargs):
        self.text(x, y, text, center=True, **kwargs)


def architecture():
    f = Figure('fig2-architecture', 224)
    f.text(98, 9, '(b) Evidence sources', bold=True, color=MUTED)
    cards = [(98, 80, BLUE, BLUE_PALE, 'Analyzer-native', 'CFG / call graph', 'Optional traces'),
             (196, 80, MUTED, PALE, 'Source context', 'Patch / source', 'Bounded context'),
             (294, 84, AMBER, AMBER_PALE, 'Analyzer output', 'Build diagnostics', 'Paired warnings')]
    for x, w, c, bg, title, line1, line2 in cards:
        f.box(x, 19, w, 43, fill=bg, stroke=c)
        f.centered(x+w/2, 32, title, size=9, bold=True, color=c, maxw=w-8)
        f.centered(x+w/2, 45, line1, size=8.3, maxw=w-8)
        f.centered(x+w/2, 56, line2, size=8.3, maxw=w-8)
        f.route([(x+w/2, 62), (x+w/2, 75)], color=LINE, arrow=False, lw=.8)
    f.route([(138,75),(336,75)],color=LINE,arrow=False,lw=.8)
    f.route([(138,75),(138,102)],color=MUTED,lw=.9)
    f.text(0, 88, '(a) Inputs', bold=True, color=MUTED)
    f.text(196, 88, '(c) Refinement', bold=True, color=MUTED)
    f.box(0,102,80,44,fill=PALE)
    f.centered(40,115,'Checker + patch',size=9,bold=True,maxw=72)
    f.centered(40,128,'Frozen source',maxw=72)
    f.centered(40,139,'V/F object scope',maxw=72)
    f.box(98,102,80,44,stroke=MUTED)
    f.centered(138,115,'Evidence bundle',size=9,bold=True,maxw=72)
    f.centered(138,128,'origin + scope',maxw=72)
    f.centered(138,139,'raw-output hash',maxw=72)
    f.box(196,102,80,44,stroke=INK)
    f.centered(236,119,'LLM edit',size=10,bold=True,maxw=72)
    f.centered(236,134,'Latest attempt L',maxw=72)
    f.box(294,102,84,44,stroke=AMBER,fill=AMBER_PALE)
    f.centered(336,119,'Compile + scan',size=9.4,bold=True,maxw=76)
    f.centered(336,134,'vulnerable / fixed',size=8.3,maxw=76)
    for a,b in [(80,98),(178,196),(276,294)]:
        f.route([(a,124),(b,124)])
    f.route([(310,146),(310,165),(236,165),(236,146)],color=AMBER,dash=True)
    f.box(242,159,62,12,stroke=None,radius=0)
    f.centered(273,168,'code + feedback',size=8,color=AMBER,maxw=62)
    f.route([(338,146),(338,189)],color=GREEN)
    f.text(344,169,'accept',size=8,color=GREEN,maxw=33)
    f.box(294,189,84,31,stroke=GREEN,fill=GREEN_PALE)
    f.centered(336,202,'Retained D*',size=9.6,bold=True,color=GREEN,maxw=76)
    f.centered(336,214,'Best accepted checker',size=7.6,maxw=77)
    f.text(0,180,'Retain only after a healthy run',size=9,bold=True)
    f.text(0,194,'Recover a vulnerable warning, or reduce fixed reports',size=8.3,maxw=282)
    f.text(0,206,'while preserving vulnerable-side warning presence.',size=8.3,maxw=282)
    return f


def study():
    f = Figure('study-design', 288)
    f.text(0,10,'Main comparison · E1',size=10,bold=True)
    f.box(0,19,378,35,fill=PALE)
    f.text(13,45,'39',size=24,bold=True,color=BLUE)
    f.text(57,33,'Generated-only CSA checkers',size=10,bold=True)
    f.text(57,46,'Same checker, patch and object scope',size=9,color=MUTED)
    f.route([(189,54),(189,63)],color=MUTED,arrow=False,lw=.9)
    f.route([(57,63),(321,63)],color=MUTED,arrow=False,lw=.9)
    for x in [57,189,321]: f.route([(x,63),(x,74)],color=MUTED,lw=.9)
    entries=[(0,'KNighter','Report-driven loop',MUTED,WHITE),
             (132,'SemWeaver','Native evidence',BLUE,BLUE_PALE),
             (264,'SemWeaver','No internal evidence',MUTED,WHITE)]
    for x,a,b,c,bg in entries:
        f.box(x,74,114,31,fill=bg,stroke=c)
        f.centered(x+57,87,a,size=9.6,bold=True,color=c,maxw=106)
        f.centered(x+57,99,b,size=8.5,maxw=106)
    f.centered(189,118,'GPT-6-Luna high · 32-reply ceiling*',size=9,color=MUTED)
    f.text(0,139,'Original 12-subject subset',size=10,bold=True)
    f.text(273,139,'Same inputs for E3 / E4',size=8,color=MUTED,maxw=105)
    f.route([(0,145),(378,145)],color=LINE,arrow=False,lw=.65)
    f.box(0,157,181,79)
    f.box(197,157,181,79)
    f.text(10,172,'Repeated ablation · E3',size=9.4,bold=True,maxw=161)
    f.text(10,185,'Native vs. no-internal',size=8.5,color=MUTED)
    for x,txt,bg,c in [(10,'1: reuse',PALE,MUTED),(65,'2: new',BLUE_PALE,BLUE),(120,'3: new',BLUE_PALE,BLUE)]:
        f.box(x,195,51,22,fill=bg,stroke=None,radius=2)
        f.centered(x+25.5,209,txt,size=8.5,color=c,maxw=45)
    f.text(10,230,'72 cells: 24 reused + 48 new',size=8.8,bold=True,maxw=161)
    f.text(207,172,'Model sensitivity · E4',size=9.4,bold=True,maxw=161)
    f.text(207,185,'Native only · 2-reply cap',size=8.5,color=MUTED)
    for x,lines in [(207,['GPT-6-Luna']),(262,['GLM-5.3']),(317,['GLM-5.3','Flash'])]:
        f.box(x,195,51,22,fill=PALE,stroke=None,radius=2)
        for i,line in enumerate(lines):
            f.centered(x+25.5,209 if len(lines)==1 else 204+i*9,line,size=7.6,maxw=47)
    f.text(207,230,'36 cells: 12 inputs × 3 models',size=8.8,bold=True,maxw=161)
    f.box(0,253,378,32,fill=PALE,stroke=LINE,dash=True)
    f.text(11,267,'CodeQL · E2',size=10,bold=True)
    f.text(225,267,'1 query × 3 decodes',size=10,bold=True,maxw=142)
    f.text(11,279,'Separate case; not pooled with the CSA comparison',size=8.5,color=MUTED)
    return f


def arrow_triangle(points, length=4, half=1.8):
    (x0,y0),(x1,y1)=points[-2:]
    distance=math.hypot(x1-x0,y1-y0)
    ux,uy=(x1-x0)/distance,(y1-y0)/distance
    return [(x1,y1),(x1-length*ux-half*uy,y1-length*uy+half*ux),
            (x1-length*ux+half*uy,y1-length*uy-half*ux)]


def export_pdf(f,path):
    c=pdfcanvas.Canvas(str(path),pagesize=(f.width,f.height),invariant=1)
    c.setAuthor('');c.setCreator('Editable research figure builder');c.setTitle(f.name)
    for o in f.items:
        c.saveState()
        if o['kind']=='box':
            c.setFillColor(HexColor(o['fill']));c.setLineWidth(o['lw'])
            if o['stroke']:c.setStrokeColor(HexColor(o['stroke']))
            if o['dash']:c.setDash(3,2)
            c.roundRect(o['x'],f.height-o['y']-o['h'],o['w'],o['h'],o['radius'],
                        stroke=bool(o['stroke']),fill=1)
        elif o['kind']=='text':
            c.setFont('FigureBold' if o['bold'] else 'FigureSans',o['size'])
            c.setFillColor(HexColor(o['color']));c.drawString(o['x'],f.height-o['y'],o['value'])
        else:
            c.setLineWidth(o['lw']);c.setStrokeColor(HexColor(o['color']))
            if o['dash']:c.setDash(3,2)
            p=c.beginPath();p.moveTo(o['points'][0][0],f.height-o['points'][0][1])
            for x,y in o['points'][1:]:p.lineTo(x,f.height-y)
            c.drawPath(p)
            if o['arrow']:
                c.setDash();c.setFillColor(HexColor(o['color']))
                a=arrow_triangle(o['points']);p=c.beginPath();p.moveTo(a[0][0],f.height-a[0][1])
                for x,y in a[1:]:p.lineTo(x,f.height-y)
                p.close();c.drawPath(p,stroke=0,fill=1)
        c.restoreState()
    c.showPage();c.save()


def export_svg(f,path):
    root=ET.Element('svg',xmlns='http://www.w3.org/2000/svg',width=f'{f.width}pt',
                    height=f'{f.height}pt',viewBox=f'0 0 {f.width} {f.height}')
    for o in f.items:
        if o['kind']=='box':
            attrs={k:str(o[k]) for k in ('x','y')}
            attrs.update(width=str(o['w']),height=str(o['h']),rx=str(o['radius']),
                         fill=o['fill'],stroke=o['stroke'] or 'none')
            attrs['stroke-width']=str(o['lw'])
            if o['dash']:attrs['stroke-dasharray']='3 2'
            ET.SubElement(root,'rect',attrs)
        elif o['kind']=='text':
            e=ET.SubElement(root,'text',{'x':str(o['x']),'y':str(o['y']),
                'font-family':'Liberation Sans, Arial, sans-serif','font-size':str(o['size']),
                'font-weight':'700' if o['bold'] else '400','fill':o['color']})
            e.text=o['value']
        else:
            attrs={'points':' '.join(f'{x},{y}' for x,y in o['points']),
                   'fill':'none','stroke':o['color'],'stroke-width':str(o['lw'])}
            if o['dash']:attrs['stroke-dasharray']='3 2'
            ET.SubElement(root,'polyline',attrs)
            if o['arrow']:
                ET.SubElement(root,'polygon',points=' '.join(f'{x},{y}' for x,y in arrow_triangle(o['points'])),fill=o['color'])
    ET.ElementTree(root).write(path,encoding='utf-8',xml_declaration=True)


def export_pptx(figures,path):
    prs=Presentation();scale=2
    prs.slide_width=Pt(402*scale);prs.slide_height=Pt(312*scale)
    prs.core_properties.author='';prs.core_properties.last_modified_by=''
    prs.core_properties.title='Editable SemWeaver diagrams'
    def rgb(v):return RGBColor.from_string(v.lstrip('#'))
    def pos(x):return Pt((x+12)*scale)
    for f in figures:
        slide=prs.slides.add_slide(prs.slide_layouts[6])
        for index,o in enumerate(f.items):
            if o['kind']=='box':
                sh=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if o['radius'] else MSO_SHAPE.RECTANGLE,
                    pos(o['x']),pos(o['y']),Pt(o['w']*scale),Pt(o['h']*scale))
                if o['radius']:sh.adjustments[0]=min(.5,o['radius']/min(o['w'],o['h']))
                sh.fill.solid();sh.fill.fore_color.rgb=rgb(o['fill'])
                if o['stroke']:
                    sh.line.color.rgb=rgb(o['stroke']);sh.line.width=Pt(o['lw']*scale)
                    if o['dash']:sh.line.dash_style=MSO_LINE_DASH_STYLE.DASH
                else:sh.line.fill.background()
            elif o['kind']=='text':
                sh=slide.shapes.add_textbox(pos(o['x']),pos(o['y']-o['ascent']),
                    Pt((o['w']+3)*scale),Pt((o['size']*1.6)*scale))
                tf=sh.text_frame;tf.clear();tf.word_wrap=False;tf.vertical_anchor=MSO_ANCHOR.TOP
                tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
                p=tf.paragraphs[0];p.space_before=p.space_after=Pt(0)
                r=p.add_run();r.text=o['value'];r.font.name='Arial'
                r.font.size=Pt(o['size']*scale);r.font.bold=o['bold'];r.font.color.rgb=rgb(o['color'])
            else:
                for i,(start,end) in enumerate(zip(o['points'],o['points'][1:])):
                    sh=slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,pos(start[0]),pos(start[1]),pos(end[0]),pos(end[1]))
                    sh.line.color.rgb=rgb(o['color']);sh.line.width=Pt(o['lw']*scale)
                    if o['dash']:sh.line.dash_style=MSO_LINE_DASH_STYLE.DASH
                    if o['arrow'] and i==len(o['points'])-2:
                        arrow=OxmlElement('a:tailEnd');arrow.set('type','triangle');arrow.set('w','sm');arrow.set('len','sm')
                        sh.line._get_or_add_ln().append(arrow)
            sh.name=f'{f.name}-{index:03d}-{o["kind"]}'
    prs.save(path)


def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--out',type=Path,default=Path(__file__).parent)
    parser.add_argument('--font-dir',type=Path,default=Path('/usr/share/fonts/truetype/liberation'))
    args=parser.parse_args();args.out.mkdir(parents=True,exist_ok=True)
    for name,suffix in [('FigureSans','Regular'),('FigureBold','Bold')]:
        pdfmetrics.registerFont(TTFont(name,str(args.font_dir/f'LiberationSans-{suffix}.ttf')))
    figures=[architecture(),study()]
    for f in figures:
        export_pdf(f,args.out/(f.name+'.pdf'));export_svg(f,args.out/(f.name+'.svg'))
        (args.out/(f.name+'.json')).write_text(json.dumps(vars(f),indent=2,ensure_ascii=False)+'\n')
    export_pptx(figures,args.out/'editable-figures.pptx')
    print(json.dumps({'figures':[f.name for f in figures],'editable_slides':len(figures)}))


if __name__=='__main__':main()
